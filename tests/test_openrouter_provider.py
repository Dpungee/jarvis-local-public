"""OpenRouter as a model provider for JARVIS and the Agent Hub (no network in these tests)."""
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from jarvis import openrouter
from jarvis.model_client import ModelClient, ModelProviderError, split_model_reference

KEY = "sk-or-v1-" + "a1b2c3d4" * 8
CATALOG = {"data": [
    {"id": "stealth/space-bunny-alpha", "name": "Space Bunny Alpha", "created": 3, "context_length": 1_000_000,
     "pricing": {"prompt": "0", "completion": "0"}, "supported_parameters": ["tools", "reasoning_effort"],
     "architecture": {"input_modalities": ["text", "image"]}},
    {"id": "vendor/paid-model", "name": "Paid", "created": 2, "context_length": 200_000,
     "pricing": {"prompt": "0.000001", "completion": "0.000002"}, "supported_parameters": ["tools"],
     "architecture": {"input_modalities": ["text"]}},
    {"id": "vendor/no-tools:free", "name": "No tools", "created": 1, "context_length": 8_000,
     "pricing": {"prompt": "0", "completion": "0"}, "supported_parameters": [],
     "architecture": {"input_modalities": ["text"]}},
]}


def load_catalog():
    openrouter._catalog.update(at=0.0, tried=0.0, models={})
    return openrouter.catalog(refresh=True, fetch=lambda url: CATALOG)


class _Response(io.BytesIO):
    headers: dict = {}

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


class FakeTransport:
    def __init__(self, *bodies):
        self.bodies = list(bodies)
        self.requests = []

    def __call__(self, request, timeout=None):
        self.requests.append((request, json.loads(request.data)))
        return _Response(json.dumps(self.bodies.pop(0)).encode())


def client(*bodies):
    transport = FakeTransport(*bodies)
    return openrouter.OpenRouterClient(KEY, open_url=transport, max_retries=0), transport


class MessageConversionTests(unittest.TestCase):
    def test_tool_calls_and_results_are_paired_by_id(self):
        converted = openrouter.chat_messages([
            {"role": "system", "content": "be useful"},
            {"role": "user", "content": [{"type": "text", "text": "what is this"},
                                         {"type": "image", "mime": "image/png", "data": "iVBOR"}]},
            {"role": "assistant", "content": "", "tool_calls": [
                {"function": {"name": "web_search", "arguments": {"query": "x"}}},
                {"function": {"name": "read_file", "arguments": "{\"path\": \"a.txt\"}"}}]},
            {"role": "tool", "content": "results"},
            {"role": "tool", "content": "file text"},
        ])
        self.assertEqual(converted[1]["content"][1],
                         {"type": "image_url", "image_url": {"url": "data:image/png;base64,iVBOR"}})
        calls = converted[2]["tool_calls"]
        self.assertIsNone(converted[2]["content"])
        self.assertEqual([c["function"]["name"] for c in calls], ["web_search", "read_file"])
        self.assertEqual([m["tool_call_id"] for m in converted[3:]], [calls[0]["id"], calls[1]["id"]])

    def test_unmatched_tool_result_is_refused(self):
        with self.assertRaises(ModelProviderError):
            openrouter.chat_messages([{"role": "tool", "content": "x"}])


class ClientTests(unittest.TestCase):
    def test_tool_call_round_trip_and_request_shape(self):
        c, transport = client({"model": "stealth/space-bunny-alpha", "choices": [{"finish_reason": "tool_calls",
                               "message": {"content": None, "tool_calls": [{"id": "x", "type": "function",
                                           "function": {"name": "web_search", "arguments": "{\"query\":\"sol\"}"}}]}}],
                               "usage": {"prompt_tokens": 1200, "completion_tokens": 30}})
        tools = [{"type": "function", "function": {"name": "web_search", "description": "search",
                                                   "parameters": {"type": "object", "properties": {}}}}]
        response = c.chat([{"role": "user", "content": "search sol"}], tools, "stealth/space-bunny-alpha", think=True)
        request, payload = transport.requests[0]
        self.assertEqual(request.full_url, openrouter.CHAT_URL)
        self.assertEqual(request.get_header("Authorization"), f"Bearer {KEY}")
        self.assertEqual(payload["model"], "stealth/space-bunny-alpha")
        self.assertEqual(payload["tools"][0]["function"]["name"], "web_search")
        self.assertEqual(payload["reasoning"], {"effort": "medium", "exclude": True})
        self.assertEqual(response["tool_calls"][0]["function"]["name"], "web_search")
        self.assertEqual(response.done_reason, "tool_use")
        self.assertTrue(response.model_attested)
        self.assertEqual(c.call_usage[-1]["context_tokens"], 1200)

    def test_a_pinned_effort_wins_and_json_output_is_requested(self):
        c, transport = client({"choices": [{"message": {"content": "{}"}, "finish_reason": "stop"}]})
        c.set_fixed_effort("high")
        c.chat([{"role": "user", "content": "x"}], [], "vendor/m", think=False, response_format="json")
        payload = transport.requests[0][1]
        self.assertEqual(payload["reasoning"]["effort"], "high")
        self.assertEqual(payload["response_format"], {"type": "json_object"})
        with self.assertRaises(ValueError):
            c.set_fixed_effort("ultracode")

    def test_quick_turns_default_to_low_effort(self):
        c, transport = client(*[{"choices": [{"message": {"content": "hi"}, "finish_reason": "stop"}]}] * 3)
        for think in (None, False, True):
            c.chat([{"role": "user", "content": "x"}], [], "vendor/m", think=think)
        self.assertEqual([r[1]["reasoning"]["effort"] for r in transport.requests], ["low", "low", "medium"])

    def test_strict_schema_failure_retries_as_plain_json(self):
        schema = {"type": "object", "properties": {"answer": {"type": "string"}}, "required": ["answer"]}
        c, transport = client({"error": {"code": 502, "message": "JSON error injected into SSE stream"}},
                              {"choices": [{"message": {"content": "{\"answer\": \"hi\"}"}, "finish_reason": "stop"}]})
        reply = c.chat([{"role": "user", "content": "x"}], [], "vendor/m", response_format=schema)
        self.assertEqual(reply["content"], '{"answer": "hi"}')
        first, second = transport.requests[0][1], transport.requests[1][1]
        self.assertEqual(first["response_format"]["type"], "json_schema")
        self.assertEqual(second["response_format"], {"type": "json_object"})
        self.assertIn("JSON schema", second["messages"][0]["content"])
        c2, _ = client({"error": {"code": 401, "message": "bad key"}})
        with self.assertRaises(ModelProviderError):
            c2.chat([{"role": "user", "content": "x"}], [], "vendor/m", response_format=schema)

    def test_an_error_body_is_a_provider_error_with_its_status(self):
        c, _ = client({"error": {"code": 402, "message": "Insufficient credits"}})
        with self.assertRaises(ModelProviderError) as caught:
            c.chat([{"role": "user", "content": "x"}], [], "vendor/m")
        self.assertEqual(caught.exception.status_code, 402)

    def test_model_client_routes_openrouter_references(self):
        c, transport = client({"choices": [{"message": {"content": "OK"}, "finish_reason": "stop"}]})
        self.assertEqual(split_model_reference("openrouter:stealth/space-bunny-alpha"),
                         ("openrouter", "stealth/space-bunny-alpha"))
        routed = ModelClient(None, openrouter=c)
        self.assertEqual(routed.chat([{"role": "user", "content": "hi"}], [],
                                     "openrouter:stealth/space-bunny-alpha")["content"], "OK")
        with self.assertRaises(ModelProviderError):
            ModelClient(None).chat([{"role": "user", "content": "hi"}], [], "openrouter:vendor/m")

    def test_the_key_alone_never_adds_a_fallback_provider(self):
        from types import SimpleNamespace
        from jarvis.model_client import build_model_client
        base = dict(data_dir=Path(tempfile.gettempdir()), ollama_enabled=False, ollama_url="http://127.0.0.1:11434",
                    ollama_allow_remote=False, ollama_health_timeout=1.0, ollama_generation_timeout=1.0,
                    ollama_max_output_tokens=256, ollama_max_response_bytes=1024, ollama_max_retries=0,
                    ollama_retry_backoff=0.1)
        with mock.patch.dict(os.environ, {"OPENROUTER_API_KEY": KEY}):
            plain = build_model_client(SimpleNamespace(model="claude-cli:x", **base))
            named = build_model_client(SimpleNamespace(model="openrouter:vendor/m", **base))
        self.assertIsNone(plain.openrouter)
        self.assertIsNotNone(named.openrouter)


class CatalogTests(unittest.TestCase):
    def tearDown(self):
        openrouter._catalog.update(at=0.0, tried=0.0, models={})

    def test_agent_models_need_tools_and_put_featured_then_free_first(self):
        load_catalog()
        self.assertEqual([m["id"] for m in openrouter.agent_models()],
                         ["stealth/space-bunny-alpha", "vendor/paid-model"])
        facts = openrouter.model_facts()["stealth/space-bunny-alpha"]
        self.assertEqual((facts["free"], facts["vision"], facts["context_length"]), (True, True, 1_000_000))
        self.assertFalse(openrouter.model_facts()["vendor/paid-model"]["free"])

    def test_cached_reads_never_fetch_and_a_failed_fetch_backs_off(self):
        calls = []

        def failing(url):
            calls.append(url)
            raise OSError("offline")
        self.assertEqual(openrouter.catalog(cached_only=True, fetch=failing), {})
        self.assertEqual(calls, [])
        openrouter.catalog(fetch=failing)
        openrouter.catalog(fetch=failing)
        self.assertEqual(len(calls), 1)

    def test_model_ids(self):
        for good in ("stealth/space-bunny-alpha", "qwen/qwen3.8-27b:free", "~openai/gpt-sol-latest"):
            self.assertTrue(openrouter.valid_model_id(good), good)
        for bad in ("space-bunny", "a/b c", "../etc/passwd", "vendor/", "x" * 70 + "/m"):
            self.assertFalse(openrouter.valid_model_id(bad), bad)


class KeyStoreTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.store = openrouter.KeyStore(Path(tmp.name))

    def test_set_get_clear_and_reject_non_keys(self):
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("OPENROUTER_API_KEY", None)
            self.assertIsNone(self.store.get())
            with self.assertRaises(ValueError):
                self.store.set("hunter2")
            self.store.set(KEY)
            self.assertEqual(self.store.get(), KEY)
            self.assertEqual(self.store.source(), "hub")
            self.store.clear()
            self.assertIsNone(self.store.get())

    def test_environment_key_is_a_fallback(self):
        with mock.patch.dict(os.environ, {"OPENROUTER_API_KEY": KEY}):
            self.assertEqual(self.store.get(), KEY)
            self.assertEqual(self.store.source(), "environment")

    def test_the_key_check_never_returns_the_key(self):
        outcome = openrouter.check_key(KEY, fetch=lambda url, headers: {"data": {
            "label": KEY[:14] + "...", "is_free_tier": True, "limit": None}})
        self.assertTrue(outcome["ok"])
        self.assertNotIn(KEY[:14], json.dumps(outcome))


class HubOpenRouterTests(unittest.TestCase):
    def setUp(self):
        from jarvis.agent_hub import HubService
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.service = HubService(state_dir=self.root, provider_profile_dir=self.root / "profile",
                                  runtime_kwargs={"autostart": False, "provider_probe": lambda _: {
                                      "installed": True, "authenticated": True, "detail": "t", "version": "t"}})
        self.addCleanup(self.service.close)
        load_catalog()
        self.addCleanup(lambda: openrouter._catalog.update(at=0.0, tried=0.0, models={}))

    def test_an_agent_can_be_created_on_an_openrouter_model(self):
        agent = self.service.create_agent({"name": "Bunny", "role": "r", "provider": "openrouter",
                                           "model": "stealth/space-bunny-alpha", "enable": True})
        self.assertEqual((agent["provider"], agent["model"]), ("openrouter", "stealth/space-bunny-alpha"))
        overview = self.service.overview()
        self.assertIn("openrouter", overview["providers"])
        self.assertEqual(overview["known_models"]["openrouter"][0], "stealth/space-bunny-alpha")
        self.assertTrue(overview["openrouter_models"]["stealth/space-bunny-alpha"]["free"])
        runtime = self.service.runtime
        self.assertEqual(runtime.model_window("openrouter", "stealth/space-bunny-alpha")["tokens"], 1_000_000)
        self.assertEqual(runtime.effort_levels("openrouter", "stealth/space-bunny-alpha"), list(openrouter.EFFORTS))

    def test_bad_models_are_refused(self):
        from jarvis.agent_hub_runtime import TaskError, validate_model
        with self.assertRaises(TaskError):
            validate_model("openrouter", "not a model")
        with self.assertRaises(TaskError):
            validate_model("openrouter-ish", "stealth/space-bunny-alpha")

    def test_the_key_is_checked_stored_and_never_returned(self):
        runtime = self.service.runtime
        runtime._provider_probe = runtime._probe_provider
        with mock.patch.object(openrouter, "check_key", return_value={"ok": True, "free_tier": True}), \
                mock.patch.object(openrouter, "refresh_catalog_in_background"):
            providers = runtime.set_openrouter_key(KEY)
            self.assertTrue(providers["openrouter"]["authenticated"])
            self.assertNotIn(KEY, json.dumps(self.service.overview(), default=str))
            self.assertEqual(runtime.openrouter_keys().get(), KEY)
            providers = runtime.clear_openrouter_key()
        self.assertFalse(providers["openrouter"]["authenticated"])
        with mock.patch.object(openrouter, "check_key", return_value={"ok": False, "error": "OpenRouter rejected the key."}):
            from jarvis.agent_hub_runtime import TaskError
            with self.assertRaises(TaskError):
                runtime.set_openrouter_key(KEY)
        self.assertIsNone(runtime.openrouter_keys().get())

    def test_the_agent_client_is_bound_to_openrouter_only(self):
        from jarvis.agent_providers import build_agent_client
        with self.assertRaises(ModelProviderError):
            build_agent_client(self.root / "profile", "openrouter", "stealth/space-bunny-alpha")
        openrouter.KeyStore(self.root / "profile").set(KEY)
        client = build_agent_client(self.root / "profile", "openrouter", "stealth/space-bunny-alpha", effort="low")
        self.assertIsNotNone(client.openrouter)
        self.assertIsNone(client.claude_cli)
        self.assertIsNone(client.codex_cli)
        self.assertEqual(client.openrouter.fixed_effort, "low")

    def test_the_agents_own_words_are_not_read_as_provider_limits(self):
        from jarvis.agent_hub_runtime import classify_failure
        state, _ = classify_failure("", "Built a capacity-based LRU cache; the quota check and login page work. "
                                        "Incomplete: tests were not rerun.")
        self.assertEqual(state, "FAILED")
        self.assertEqual(classify_failure("", "Jarvis could not reach a usable model: rate limit")[0], "WAITING_PROVIDER")

    def test_hub_agents_get_room_for_multi_step_jobs(self):
        from jarvis.agent_hub_runtime import DEFAULT_PERMISSIONS
        config = self.service.runtime._agent_config(agent_id="a", project_root=self.root, permissions=DEFAULT_PERMISSIONS,
                                                    reference="openrouter:stealth/space-bunny-alpha")
        self.assertGreaterEqual(config.max_steps, 40)

    def test_coding_on_a_chosen_model_gets_the_coding_allowance(self):
        import dataclasses
        from types import SimpleNamespace
        from jarvis.agent import Agent
        from jarvis.config import Config
        from jarvis.memory import Memory
        (self.root / "ws").mkdir()
        config = dataclasses.replace(Config.load(), workspace=self.root / "ws", data_dir=self.root,
                                     ollama_enabled=False, claude_cli_enabled=False, codex_cli_enabled=False,
                                     max_steps=40)
        memory = Memory(self.root / "budget.db")
        self.addCleanup(memory.close)
        agent = Agent(config, memory, client=SimpleNamespace(models=lambda refresh=True: ["x"], chat=None),
                      record_training=False)
        route = SimpleNamespace(profile="custom", model="openrouter:vendor/m")
        budgets = dict(staged_tool_calls=0, learning_task=False, skill_authoring_task=False)
        self.assertEqual(agent._phase_tool_budgets(route, requires_coding=True, **budgets), (12, 16))
        agent.open_toolset = True  # a Hub personal agent: no hard ceiling, progress-gated
        self.assertEqual(agent._phase_tool_budgets(route, requires_coding=True, **budgets),
                         (28, Agent._HUB_OPEN_ENDED_BUDGET))
        self.assertEqual(agent._phase_tool_budgets(route, requires_coding=False, **budgets), (12, 16))

    def test_key_and_credit_failures_wait_for_the_operator(self):
        from jarvis.agent_hub_runtime import classify_failure
        self.assertEqual(classify_failure("OpenRouter model provider request failed with HTTP 401", "")[0],
                         "WAITING_PROVIDER")
        self.assertIn("credits", classify_failure("", "OpenRouter model provider request failed with HTTP 402")[1])


class SubagentTests(unittest.TestCase):
    def setUp(self):
        import threading
        from types import SimpleNamespace
        from jarvis.agent_hub import HubService
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.seen, self.lock = [], threading.Lock()
        self.running, self.peak = 0, 0

        def factory(config, memory, on_event, client):
            test = self

            class FakeAgent:
                def __init__(self):
                    self.toolbox = SimpleNamespace(tools={"web_search": 1, "read_file": 1, "run_process": 1,
                                                          "browser_open": 1, "spawn_subagents": 1},
                                                   execute=lambda name, arguments: '{"ok": true, "result": {}}')
                    self.operator_brief = ""
                    self.open_toolset = False

                def run(self, text, conversation_id=None, cancellation_guard=None, **_):
                    import time as clock
                    with test.lock:
                        test.running += 1
                        test.peak = max(test.peak, test.running)
                        test.seen.append((text, sorted(self.toolbox.tools), self.operator_brief, self.open_toolset))
                    self.toolbox.execute("web_search", {"query": text})
                    clock.sleep(0.3)
                    with test.lock:
                        test.running -= 1
                    if "explode" in text:
                        raise RuntimeError("boom")
                    return f"report on {text}"
            return FakeAgent()

        self.service = HubService(state_dir=self.root, provider_profile_dir=self.root / "profile",
                                  runtime_kwargs={"autostart": False, "agent_factory": factory,
                                                  "client_factory": lambda config: SimpleNamespace(),
                                                  "provider_probe": lambda _: {"installed": True, "authenticated": True,
                                                                               "detail": "t", "version": "t"}})
        self.addCleanup(self.service.close)

    def spawn(self, helpers, minutes=1):
        import threading
        from types import SimpleNamespace
        agent = self.service.create_agent({"name": "Lead", "role": "r", "provider": "claude-cli",
                                           "model": "claude-opus-5-5", "enable": True})
        chat = self.service.create_chat(agent["agent_id"], {"title": "t"})["chat_id"]
        task = self.service.chat(agent["agent_id"], chat, {"body": "research", "request_id": "s1"})
        full = dict(task, agent_id=agent["agent_id"], project_id=task.get("project_id") or "command-center")
        return self.service.runtime._spawn_subagents(full, SimpleNamespace(display_name="Lead"),
                                                     "claude-cli:claude-opus-5-5", self.root,
                                                     {"cancel": threading.Event()}, helpers, minutes)

    def test_helpers_run_in_parallel_with_research_tools_only(self):
        out = self.spawn([{"name": "launched", "task": "coins launched on Robinhood Chain"},
                          {"name": "gaps", "task": "ideas not launched yet"}])
        self.assertEqual([h["status"] for h in out["helpers"]], ["done", "done"])
        self.assertEqual(out["helpers"][0]["report"], "report on coins launched on Robinhood Chain")
        self.assertEqual(out["helpers"][0]["tool_calls"], 1)
        self.assertEqual(self.peak, 2)
        for _text, tools, brief, open_toolset in self.seen:
            self.assertEqual(tools, ["read_file", "web_search"])
            self.assertIn("helper agent", brief)
            self.assertTrue(open_toolset)
        events = [e["summary"] for e in self.service.runtime.latest_events(limit=50)]
        self.assertTrue(any(s.startswith("launched · ") for s in events))

    def test_one_failing_helper_does_not_sink_the_others(self):
        out = self.spawn([{"name": "a", "task": "fine"}, {"name": "b", "task": "explode"}])
        self.assertEqual([h["status"] for h in out["helpers"]], ["done", "failed"])

    def test_no_cap_on_helpers_all_run_at_once(self):
        out = self.spawn([{"name": f"h{i}", "task": f"part {i}"} for i in range(6)])
        self.assertEqual([h["status"] for h in out["helpers"]], ["done"] * 6)
        self.assertEqual(self.peak, 6)
        from jarvis.agent_hub_runtime import SUBAGENT_PARAMETERS
        self.assertNotIn("maxItems", SUBAGENT_PARAMETERS["properties"]["helpers"])

    def test_limits(self):
        with self.assertRaises(ValueError):
            self.spawn([])
        with self.assertRaises(ValueError):
            self.spawn([{"name": "a", "task": ""}])

    def test_the_permission_offers_the_tool_and_the_brief_explains_it(self):
        from types import SimpleNamespace as NS
        from jarvis.agent_hub_runtime import allowed_tools
        self.assertIn("spawn_subagents", allowed_tools({"subagents": True}))
        brief = self.service.runtime._operator_brief({"attempt": 1}, NS(display_name="R", role="", purpose=""), "",
                                                     {"subagents": True})
        self.assertIn("spawn_subagents", brief)


PNG = bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                    "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")


class ImageMessageTests(unittest.TestCase):
    def setUp(self):
        from jarvis.agent_hub import HubService
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.service = HubService(state_dir=self.root, provider_profile_dir=self.root / "profile",
                                  runtime_kwargs={"autostart": False, "provider_probe": lambda _: {
                                      "installed": True, "authenticated": True, "detail": "t", "version": "t"}})
        self.addCleanup(self.service.close)
        agent = self.service.create_agent({"name": "Eyes", "role": "r", "provider": "claude-cli",
                                           "model": "claude-opus-5-5", "enable": True})
        self.agent_id = agent["agent_id"]
        self.chat_id = self.service.create_chat(self.agent_id, {"title": "t"})["chat_id"]

    def send(self, images, request_id="i1"):
        import base64
        return self.service.chat(self.agent_id, self.chat_id, {
            "body": "what does this screenshot say", "request_id": request_id,
            "images": [{"name": n, "mime": m, "data": base64.b64encode(d).decode()} for n, m, d in images]})

    def test_images_are_kept_for_the_run_and_named_in_the_message(self):
        task = self.send([("screen.png", "image/png", PNG)])
        self.assertIn("📎 Attached: screen.png", task["request"])
        loaded = self.service.runtime._load_attachments(task["task_id"])
        self.assertEqual([(a.name, a.mime, a.data) for a in loaded], [("screen.png", "image/png", PNG)])

    def test_bad_images_are_refused(self):
        from jarvis.agent_hub import HubError
        with self.assertRaises(HubError):
            self.send([("fake.png", "image/png", b"not an image")])
        with self.assertRaises(HubError):
            self.send([(f"{i}.png", "image/png", PNG) for i in range(5)], "i2")

    def test_messages_without_images_are_unchanged(self):
        task = self.service.chat(self.agent_id, self.chat_id, {"body": "hi", "request_id": "i3"})
        self.assertEqual(task["request"], "hi")
        self.assertEqual(self.service.runtime._load_attachments(task["task_id"]), ())


class OpenTurnImageTests(unittest.TestCase):
    def test_personal_agent_turns_show_the_model_the_images(self):
        import dataclasses
        from types import SimpleNamespace
        from jarvis.agent import Agent
        from jarvis.attachments import ImageAttachment
        from jarvis.config import Config
        from jarvis.memory import Memory
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        (root / "ws").mkdir()
        (root / "data").mkdir()
        config = dataclasses.replace(Config.load(), workspace=root / "ws", data_dir=root / "data",
                                     ollama_enabled=False, claude_cli_enabled=False, codex_cli_enabled=False)
        memory = Memory(root / "data" / "jarvis.db")
        self.addCleanup(memory.close)
        agent = Agent(config, memory, client=SimpleNamespace(models=lambda refresh=True: ["claude-cli:x"],
                                                            chat=lambda *a, **k: None), record_training=False)
        seen = []

        thinking = []

        def fake_chat(messages, tools, route, **kwargs):
            seen.append(list(messages))
            thinking.append(kwargs.get("think_override"))
            return {"role": "assistant", "content": "It says hello."}, route
        agent._chat = fake_chat
        conversation = memory.new_conversation("t")
        with agent.toolbox.approval_context(f"conversation:{conversation}"):
            agent._open_agent_turn(conversation_id=conversation, operator_prompt="read this",
                                   route=SimpleNamespace(model="claude-cli:x", profile="general", reason="x"),
                                   recent_messages=[], attachments=(ImageAttachment("image/png", PNG, "s.png"),))
        system = seen[0][0]["content"]
        # remember is offered here, so the prompt tells the agent to save with it.
        self.assertIn("call remember", system)
        self.assertNotIn("cannot store, update, or forget durable facts", system)
        self.assertIn("cannot store, update, or forget durable facts",
                      agent.system_prompt("x", include_memory=False))
        content = seen[0][-1]["content"]
        self.assertIsInstance(content, list)
        self.assertEqual(thinking, [None])  # a job keeps the route's thinking
        with agent.toolbox.approval_context(f"conversation:{conversation}"):
            agent._open_agent_turn(conversation_id=conversation, operator_prompt="hey whats up",
                                   route=SimpleNamespace(model="claude-cli:x", profile="reasoning", reason="x"),
                                   recent_messages=[], quick=True)
        self.assertEqual(thinking[-1], False)  # small talk answers quickly
        self.assertIn("untrusted_image_attachments", content[0]["text"])
        self.assertEqual(content[1]["type"], "image")


class ReadDocumentTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)

    def test_word_and_excel_files_are_read(self):
        import docx
        import openpyxl
        from jarvis.office_reader import read_document
        document = docx.Document()
        document.add_paragraph("Quarterly budget memo")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "rent", "1200"
        document.save(self.root / "memo.docx")
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Costs"
        sheet.append(["item", "usd"])
        sheet.append(["rent", 1200])
        sheet["C2"] = "=B2*2"
        book.save(self.root / "costs.xlsx")
        word = read_document(self.root, "memo.docx")
        self.assertIn("Quarterly budget memo", word["content"])
        self.assertIn("rent\t1200", word["content"])
        excel = read_document(self.root, "costs.xlsx")
        self.assertIn("[sheet Costs]", excel["content"])
        self.assertIn("rent\t1200", excel["content"])
        self.assertEqual(excel["sheets"][0]["name"], "Costs")

    def test_paths_stay_in_the_project_and_text_files_use_read_file(self):
        from jarvis.office_reader import read_document
        (self.root / "notes.txt").write_text("x", encoding="utf-8")
        with self.assertRaises(ValueError):
            read_document(self.root, "notes.txt")
        with self.assertRaises(Exception):
            read_document(self.root, "../outside.docx")

    def test_the_hub_offers_it_with_file_reading(self):
        from jarvis.agent_hub_runtime import allowed_tools
        self.assertIn("read_document", allowed_tools({"files_read": True}))


class LauncherAndDocumentTargetTests(unittest.TestCase):
    def test_one_app_listed_twice_launches_its_start_menu_entry(self):
        from jarvis.windows_apps import InstalledApplication, WindowsAppController
        controller = WindowsAppController(Path(tempfile.gettempdir()), Path(tempfile.mkdtemp()))
        paths_entry = InstalledApplication("Notepad", Path("C:/x/Notepad.exe"), "app-paths")
        start_entry = InstalledApplication("Notepad", None, "start-apps", "Microsoft.WindowsNotepad_8wekyb3d8bbwe!App")
        other = InstalledApplication("Notepad++", Path("C:/y/notepad++.exe"), "app-paths")
        controller.catalog = lambda: [paths_entry, start_entry, other]
        self.assertIs(controller._resolve_app("notepad"), start_entry)
        self.assertIs(controller._resolve_app("Microsoft.WindowsNotepad_8wekyb3d8bbwe!App"), start_entry)
        controller.catalog = lambda: [paths_entry, other]
        with self.assertRaises(ValueError):
            controller._resolve_app("note")  # two different applications still ask

    def test_launch_as_a_noun_is_not_a_request_to_open_a_page(self):
        import jarvis.agent as agent_module
        research = ("Research tokens on Robinhood Chain: name, ticker, launch date, and the exact source URL "
                    "for each, using robinhood.com/arc and DefiLlama.")
        self.assertIsNone(agent_module._requested_browser_url(research))
        self.assertEqual(agent_module._requested_browser_url("open https://example.com/page"),
                         "https://example.com/page")

    def test_a_named_document_keeps_every_word_of_its_name(self):
        import jarvis.agent as agent_module
        required, _ = agent_module._required_effect_tools(
            "make me a word document called Q4 plan.docx and a sheet budget.xlsx",
            requires_coding=False, allow_external_mutation=False)
        self.assertIn("__effect_path__:q4 plan.docx", required)
        self.assertIn("__effect_path__:budget.xlsx", required)


class FollowUpRoutingTests(unittest.TestCase):
    def test_build_it_after_an_offer_goes_to_the_open_turn(self):
        import dataclasses
        from types import SimpleNamespace
        from jarvis.agent import Agent, _REFERS_BACK
        from jarvis.config import Config
        from jarvis.memory import Memory
        for phrase in ("build it", "yes do it", "go ahead", "sounds good, make it"):
            self.assertTrue(_REFERS_BACK.match(phrase), phrase)
        for phrase in ("build me a snake game", "fix the bug in main.py", "run the tests"):
            self.assertFalse(_REFERS_BACK.match(phrase), phrase)
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        (root / "ws").mkdir()
        (root / "data").mkdir()
        (root / "ws" / "paper_trader.py").write_text("print('bot')", encoding="utf-8")
        config = dataclasses.replace(Config.load(), workspace=root / "ws", data_dir=root / "data",
                                     ollama_enabled=False, claude_cli_enabled=False, codex_cli_enabled=False)
        memory = Memory(root / "data" / "jarvis.db")
        self.addCleanup(memory.close)
        events = []
        agent = Agent(config, memory, client=SimpleNamespace(models=lambda refresh=True: ["claude-cli:x"],
                                                            chat=lambda *a, **k: None),
                      on_event=events.append, record_training=False)
        agent.open_toolset = True
        agent._chat = lambda messages, tools, route, **kwargs: ({"role": "assistant", "content": "On it."}, route)
        conversation = memory.new_conversation("t")
        memory.add_message(conversation, "user", "what coin does good on robinhood chain")
        memory.add_message(conversation, "assistant", "I can write a Python script that pulls DefiLlama's API and "
                           "builds a TVL and fee ranking app for every protocol on the chain. Want me to build it?")
        from jarvis.agent import _is_contextual_software_build_request
        self.assertTrue(_is_contextual_software_build_request("build it", memory.recent_messages(conversation, limit=8)))
        agent.run("build it", conversation_id=conversation)
        self.assertIn("personal agent - follow-up to the conversation", events)
        self.assertIn("personal agent - all granted tools offered", events)


class BatteryFindingTests(unittest.TestCase):
    """Regressions found by the 100-case live battery on 2026-09-26."""

    def test_schedule_requests_are_recognised_without_catching_ordinary_questions(self):
        from jarvis.agent import _SCHEDULE_REQUEST
        for phrase in ("every day at 8am send me a quick weather summary for Chicago", "cancel the daily weather job",
                       "what recurring jobs do you have set up?", "remind me tomorrow at 9 to call mom",
                       "check bitcoin every 30 minutes and tell me if it drops", "send me a daily summary of AI news"):
            self.assertTrue(_SCHEDULE_REQUEST.search(phrase), phrase)
        for phrase in ("whats the weather in miami this weekend", "what is the daily active user count of twitter",
                       "build me a snake game"):
            self.assertFalse(_SCHEDULE_REQUEST.search(phrase), phrase)

    def test_input_files_are_not_output_targets(self):
        import jarvis.agent as agent_module

        def targets(prompt):
            required, _ = agent_module._required_effect_tools(prompt, requires_coding=False, allow_external_mutation=False)
            return sorted(x.split(":", 1)[1] for x in required if x.startswith("__effect_path__:"))
        self.assertEqual(targets("write a short report report_sales.docx summarising sales.csv"), ["report_sales.docx"])
        self.assertEqual(targets("make a chart.xlsx from data.csv"), ["chart.xlsx"])
        self.assertEqual(targets("read contract.docx and write summary.md"), ["summary.md"])
        self.assertEqual(targets("write the summary to notes.md and create report.docx"), ["notes.md", "report.docx"])

    def test_places_questions_are_not_shopping(self):
        from jarvis.agent import _product_research_request
        for phrase in ("find 3 highly rated sushi restaurants in Austin TX with links",
                       "recommend 3 hotels near me with links", "find a coffee shop nearby with links"):
            self.assertFalse(_product_research_request(phrase), phrase)
        for phrase in ("find me the best gaming laptop options under $1500", "compare golf clubs models to buy",
                       "find 3 espresso machine options for me", "which hotel booking deals should I buy"):
            self.assertTrue(_product_research_request(phrase), phrase)

    def test_play_it_means_open_it(self):
        from jarvis.agent import _LAUNCH_INTENT
        from jarvis.natural_language import operator_action_text
        for phrase in ("build a memory card matching game and let me play it", "make a game so i can play it",
                       "make me a todo app and show it to me", "build a snake game and open it"):
            self.assertTrue(_LAUNCH_INTENT.search(operator_action_text(phrase)), phrase)
        for phrase in ("write a python script i can use", "show me the code for bubble sort", "fix the bug in calc.py"):
            self.assertFalse(_LAUNCH_INTENT.search(operator_action_text(phrase)), phrase)

    def test_a_passing_browser_check_counts_as_the_launch(self):
        source = Path(__file__).resolve().parents[1].joinpath("jarvis", "agent.py").read_text(encoding="utf-8")
        block = source[source.index('name == "web_app_check"\n'):][:1500]
        self.assertIn('successful_tools.add("__artifact_launched__")', block)


if __name__ == "__main__":
    unittest.main()


class HubAgentsAreNeverCutOffTests(unittest.TestCase):
    """A Hub agent keeps working past the old 20-step / 30-tool cap until done or stalled."""

    def setUp(self):
        import dataclasses
        from types import SimpleNamespace
        from jarvis.agent import Agent
        from jarvis.config import Config
        from jarvis.memory import Memory
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        (root / "ws").mkdir()
        (root / "data").mkdir()
        config = dataclasses.replace(Config.load(), workspace=root / "ws", data_dir=root / "data",
                                     ollama_enabled=False, claude_cli_enabled=False, codex_cli_enabled=False)
        self.memory = Memory(root / "data" / "jarvis.db")
        self.addCleanup(self.memory.close)
        self.agent = Agent(config, self.memory, client=SimpleNamespace(models=lambda refresh=True: ["claude-cli:x"],
                                                                      chat=lambda *a, **k: None),
                           record_training=False)
        self.executed = []
        self.agent.toolbox.execute = lambda name, arguments: (
            self.executed.append((name, dict(arguments))) or json.dumps({"ok": True, "content": "fine"}))

    def run_turn(self, script, open_toolset=True):
        from types import SimpleNamespace
        self.agent.open_toolset = open_toolset
        self.sizes, self.final_notes = [], []

        def fake_chat(messages, tools, route, **kwargs):
            self.sizes.append(len(messages))
            if not tools:
                self.final_notes.append(messages[-1]["content"])
                return {"role": "assistant", "content": "final report"}, route
            return script(len(self.sizes), messages), route
        self.agent._chat = fake_chat
        conversation = self.memory.new_conversation("t")
        with self.agent.toolbox.approval_context(f"conversation:{conversation}"):
            return self.agent._open_agent_turn(conversation_id=conversation, operator_prompt="build the engine",
                                               route=SimpleNamespace(model="claude-cli:x", profile="custom",
                                                                     reason="x"),
                                               recent_messages=[])

    @staticmethod
    def call(path):
        return {"role": "assistant", "content": "",
                "tool_calls": [{"function": {"name": "read_file", "arguments": {"path": path}}}]}

    def test_a_long_job_runs_past_the_old_cap_to_completion(self):
        def script(n, _messages):
            return self.call(f"file{n}.py") if n <= 75 else {"role": "assistant", "content": "all done"}
        result = self.run_turn(script)
        self.assertEqual(len(self.executed), 75)  # old cap: 20 steps / 30 tool calls
        self.assertEqual(self.final_notes, [])  # finished on its own, never told to stop
        self.assertIn("all done", str(result))  # a real short answer is kept as it is
        self.assertLess(max(self.sizes), 60)  # checkpoints keep the working history bounded

    def test_checkpoints_keep_the_task_and_a_progress_note(self):
        notes = []

        def script(n, messages):
            if n == 45:
                notes.extend(m["content"] for m in messages if m["role"] == "user")
                return {"role": "assistant", "content": "done"}
            return self.call(f"f{n}.py")
        self.run_turn(script)
        self.assertTrue(any("build the engine" in str(n) for n in notes))
        self.assertTrue(any(str(n).startswith("Checkpoint - you are part-way") and "f1.py" in str(n)
                            for n in notes))

    def test_a_stalled_agent_is_stopped_and_asked_to_report(self):
        result = self.run_turn(lambda n, _m: self.call("same.py"))  # the same call forever
        self.assertLessEqual(len(self.executed), 60)
        self.assertEqual(len(self.final_notes), 1)
        self.assertIn("no new progress", self.final_notes[0])
        self.assertIn("final report", str(result))

    def test_a_placeholder_final_reply_is_replaced_by_the_full_answer(self):
        def script(n, _messages):
            if n <= 3:
                return self.call(f"page{n}.md")
            return {"role": "assistant", "content": "test"}
        from types import SimpleNamespace
        full = "Qdrant, Weaviate and Milvus compared, with sources: https://github.com/qdrant/qdrant/releases"
        self.sizes, self.final_notes = [], []
        self.agent.open_toolset = True

        def fake_chat(messages, tools, route, **kwargs):
            self.sizes.append(len(messages))
            if not tools:
                self.final_notes.append(messages[-1]["content"])
                return {"role": "assistant", "content": full}, route
            return script(len(self.sizes), messages), route
        self.agent._chat = fake_chat
        conversation = self.memory.new_conversation("t")
        with self.agent.toolbox.approval_context(f"conversation:{conversation}"):
            result = self.agent._open_agent_turn(conversation_id=conversation, operator_prompt="research",
                                                 route=SimpleNamespace(model="claude-cli:x", profile="custom",
                                                                       reason="x"), recent_messages=[])
        self.assertIn("placeholder", self.final_notes[0])
        self.assertIn("Qdrant, Weaviate and Milvus", str(result))

    def test_plain_jarvis_keeps_its_step_limit(self):
        self.run_turn(lambda n, _m: self.call(f"x{n}.py"), open_toolset=False)
        self.assertLessEqual(len(self.executed), 20)
        self.assertIn("Step limit reached", self.final_notes[0])


class SubagentTimeLimitTests(SubagentTests):
    def test_helpers_have_no_default_time_limit(self):
        from jarvis.agent_hub_runtime import SUBAGENT_PARAMETERS
        self.assertNotIn("maximum", SUBAGENT_PARAMETERS["properties"]["minutes"])
        out = self.spawn([{"name": "a", "task": "fine"}], minutes=None)
        self.assertEqual([h["status"] for h in out["helpers"]], ["done"])


class HubCodingRoutingTests(unittest.TestCase):
    """Codex-style prompts on a Hub agent reach the coding lane or the open turn, not research."""

    def route_events(self, prompt, open_toolset=True):
        import dataclasses
        from types import SimpleNamespace
        from jarvis.agent import Agent
        from jarvis.config import Config
        from jarvis.memory import Memory
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        (root / "ws").mkdir()
        (root / "data").mkdir()
        config = dataclasses.replace(Config.load(), workspace=root / "ws", data_dir=root / "data",
                                     ollama_enabled=False, claude_cli_enabled=False, codex_cli_enabled=False,
                                     max_steps=40)
        memory = Memory(root / "data" / "jarvis.db")
        self.addCleanup(memory.close)
        events = []
        agent = Agent(config, memory, client=SimpleNamespace(models=lambda refresh=True: ["openrouter:v/m"], chat=None),
                      record_training=False, on_event=events.append)
        agent.open_toolset = open_toolset
        agent._chat = lambda messages, tools, route, **kwargs: ({"role": "assistant", "content": "Done."}, route)
        agent.run(prompt, conversation_id=memory.new_conversation("t"))
        return events

    def test_a_feature_branch_is_coding_not_capability_research(self):
        prompt = "Initialise a git repository here, commit a README, then create a branch called feature/hello."
        self.assertFalse(any("researching" in e for e in self.route_events(prompt)))
        self.assertTrue(any("researching" in e for e in self.route_events(prompt, open_toolset=False)))

    def test_plain_text_files_use_the_open_turn(self):
        events = self.route_events("write hello.txt containing hi")
        self.assertIn("personal agent - all granted tools offered", events)

    def test_office_documents_and_self_improvement_keep_their_lanes(self):
        self.assertNotIn("personal agent - all granted tools offered",
                         self.route_events("Make me a Word document named brief.docx about unit tests"))
        self.assertTrue(any("researching" in e for e in self.route_events("learn new tools you are missing")))

    def test_hub_coding_harness_offers_the_hub_git_tools(self):
        import dataclasses
        from types import SimpleNamespace
        from jarvis.agent import Agent
        from jarvis.config import Config
        from jarvis.memory import Memory
        from jarvis.tools import Tool

        def offered_for(open_toolset):
            tmp = tempfile.TemporaryDirectory()
            self.addCleanup(tmp.cleanup)
            root = Path(tmp.name)
            (root / "ws").mkdir()
            (root / "data").mkdir()
            config = dataclasses.replace(Config.load(), workspace=root / "ws", data_dir=root / "data",
                                         ollama_enabled=False, claude_cli_enabled=False, codex_cli_enabled=False,
                                         max_steps=40)
            memory = Memory(root / "data" / "jarvis.db")
            self.addCleanup(memory.close)
            agent = Agent(config, memory, client=SimpleNamespace(models=lambda refresh=True: ["openrouter:v/m"],
                                                                chat=None), record_training=False)
            agent.open_toolset = open_toolset
            agent.toolbox.tools["git_init"] = Tool("git_init", "Start a Git repository.",
                                                   {"type": "object", "properties": {}}, lambda: "{}")
            seen = set()

            def fake_chat(messages, tools, route, **kwargs):
                seen.update(t["function"]["name"] for t in tools)
                return {"role": "assistant", "content": "Done."}, route
            agent._chat = fake_chat
            agent.run("Initialise a git repository in this folder and commit the README.",
                      conversation_id=memory.new_conversation("t"))
            return seen
        self.assertIn("git_init", offered_for(True))
        self.assertNotIn("git_init", offered_for(False))

    def test_hub_deep_research_is_an_iterative_open_turn(self):
        prompt = ("Deep research: compare three open-source vector databases (features, licence, latest version) "
                  "with a sourced summary table.")
        events = self.route_events(prompt)
        self.assertIn("personal agent - all granted tools offered", events)
        self.assertFalse(any("deterministic deep evidence" in e for e in events))
        self.assertNotIn("personal agent - all granted tools offered", self.route_events(prompt, open_toolset=False))

    def test_git_only_hub_jobs_are_accepted_by_the_coding_verifier(self):
        from jarvis.agent import Agent
        base = dict(content="Committed README.md on main and created feature/hello.", done_reason=None,
                    requires_web=False, requires_coding=True, learning_task=False, verified_urls=set())
        self.assertIsNone(Agent._acceptance_failure(
            successful_tools={"git_init", "git_commit", "git_branch", "run_process"}, **base))
        # A file write still needs the full reread / test / review chain.
        self.assertIsNotNone(Agent._acceptance_failure(
            successful_tools={"git_commit", "write_file", "__inspected_before_write__"}, **base))
        self.assertEqual(Agent._acceptance_failure(successful_tools={"run_process"}, **base),
                         "Coding work was not inspected before modification.")

    def test_self_capability_wording(self):
        from jarvis.agent import _is_hub_self_capability_request
        self.assertTrue(_is_hub_self_capability_request("add the missing capabilities to your own toolset"))
        self.assertFalse(_is_hub_self_capability_request("implement the login feature in app.py"))
        self.assertFalse(_is_hub_self_capability_request("build a workflow that emails me"))


class ClaudeCliStubAnswerTests(unittest.TestCase):
    """A placeholder in the structured field never replaces the model's finished answer."""

    def run_chat(self, structured_content, result_text):
        import subprocess
        from jarvis import model_client as mc
        client = mc.ClaudeCLIClient("claude.exe", working_directory=".")
        payload = {"type": "result", "is_error": False, "usage": {},
                   "structured_output": {"content": structured_content, "tool_calls": []},
                   "result": result_text}

        def fake_run(args, *, prompt, timeout, flags, cancellation_guard):
            return subprocess.CompletedProcess(args, 0, json.dumps(payload), "")
        tools = [{"type": "function", "function": {"name": "web_fetch", "description": "Fetch.",
                                                   "parameters": {"type": "object", "properties": {}}}}]
        with mock.patch.object(client, "_run_cli", side_effect=fake_run):
            return client.chat([{"role": "user", "content": "research"}], tools, "claude-sonnet-5")["content"]

    def test_stub_is_replaced_by_the_visible_answer(self):
        answer = "## Comparison\n\n| Database | Licence |\n|---|---|\n| Qdrant | Apache-2.0 |\n" + "Detail. " * 40
        self.assertEqual(self.run_chat("test", answer), answer.strip())

    def test_a_content_only_answer_is_valid_and_parses(self):
        from jarvis.model_client import _claude_cli_output_schema
        schema = _claude_cli_output_schema(["web_fetch"], None)
        self.assertNotIn("required", schema)
        self.assertNotIn("anyOf", schema)  # the API refuses anyOf at a tool schema's top level
        import subprocess
        from jarvis import model_client as mc
        client = mc.ClaudeCLIClient("claude.exe", working_directory=".")
        answer = "| A | B |\n|---|---|\n| 1 | 2 |"
        payload = {"type": "result", "is_error": False, "usage": {}, "structured_output": {"content": answer}}
        tools = [{"type": "function", "function": {"name": "web_fetch", "description": "Fetch.",
                                                   "parameters": {"type": "object", "properties": {}}}}]
        with mock.patch.object(client, "_run_cli", side_effect=lambda args, **kw: subprocess.CompletedProcess(
                args, 0, json.dumps(payload), "")):
            reply = client.chat([{"role": "user", "content": "table"}], tools, "claude-sonnet-5")
        self.assertEqual(reply["content"], answer)
        self.assertNotIn("tool_calls", reply)

    def test_leaked_tool_call_markup_is_stripped_or_used(self):
        from jarvis.model_client import _split_leaked_tool_call_markup
        answer = "BTC is about $82,937 per Coinbase."
        self.assertEqual(_split_leaked_tool_call_markup(answer + '</content>\n<parameter name="tool_calls">[]'),
                         (answer, []))
        self.assertEqual(_split_leaked_tool_call_markup(answer + "\n</content>"), (answer, []))
        self.assertEqual(_split_leaked_tool_call_markup(answer), (answer, []))
        text, calls = _split_leaked_tool_call_markup(
            'Checking.</content>\n<parameter name="tool_calls">[{"name": "web_fetch", "arguments": "{\\"url\\": \\"u\\"}"}]'
            '</parameter>')
        self.assertEqual(text, "Checking.")
        self.assertEqual(calls, [{"name": "web_fetch", "arguments": {"url": "u"}}])
        self.assertEqual(self.run_chat(answer + '</content>\n<parameter name="tool_calls">[]', ""), answer)

    def test_double_escaped_newlines_become_real_line_breaks(self):
        from jarvis.model_client import _unescape_double_escaped_text
        bs = chr(92)
        self.assertEqual(_unescape_double_escaped_text("BTC" + bs + "n" + bs + "n- a" + bs + "n- b"), "BTC\n\n- a\n- b")
        code = "line\nprint('a" + bs + "nb')"  # a real newline: literal escapes are code, keep them
        self.assertEqual(_unescape_double_escaped_text(code), code)
        single = "one " + bs + "n only"
        self.assertEqual(_unescape_double_escaped_text(single), single)

    def test_real_answers_and_json_results_are_left_alone(self):
        long_answer = "A complete answer that is clearly longer than forty characters in total."
        self.assertEqual(self.run_chat(long_answer, "something else " * 30), long_answer)
        self.assertEqual(self.run_chat("test", json.dumps({"content": "test", "tool_calls": []})), "test")
        self.assertEqual(self.run_chat("Done.", "Done. Short."), "Done.")
